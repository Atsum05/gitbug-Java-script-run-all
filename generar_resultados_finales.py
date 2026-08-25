#!/usr/bin/env python3
"""
Paso 2: Lee metrics.csv (combinado JUAMPI + MUTATION) y genera
resultados_finales.xlsx con el análisis estadístico comparativo.

Hojas del Excel:
  - Vista Global   : resumen ejecutivo con totales y promedios generales
  - Resumen        : una fila por bug (promedio de seeds) JUAMPI vs MUTATION
  - Detección      : detalle BUG-REVEALING por bug+seed+java_version
  - Estadísticos   : Wilcoxon signed-rank + Vargha-Delaney A12 por métrica

Emparejamiento: para cada métrica se comparan corridas con el MISMO
(bug_id, class_under_test, seed, java_version) entre JUAMPI y MUTATION
(diseño de pares, ya que ambos algoritmos corrieron con las mismas seeds).
Filas sin ambos algoritmos, o con el valor de la métrica vacío, se excluyen
de esa métrica puntual.

Si metrics.csv tiene corridas duplicadas (ej. por --force al re-correr),
se conserva la última fila no-error para esa combinación
(bug_id, class_under_test, algorithm, seed, java_version).

Uso:
    python3 generar_resultados_finales.py \
        --metrics ~/Tesis/gitbug-batch/metrics.csv \
        --out ~/Tesis/gitbug-batch/resultados_finales.xlsx
"""

import argparse
import csv
from collections import defaultdict
from pathlib import Path
from statistics import mean

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from scipy.stats import wilcoxon
except ImportError:
    wilcoxon = None

DEFAULT_METRICS = Path.home() / "Tesis" / "gitbug-batch" / "metrics.csv"
DEFAULT_OUT = Path.home() / "Tesis" / "gitbug-batch" / "resultados_finales.xlsx"

METRICAS = [
    ("line_coverage", "Line Coverage (%)", "mayor_es_mejor"),
    ("branch_coverage", "Branch Coverage (%)", "mayor_es_mejor"),
    ("mutation_score_final", "Mutation Score (%)", "mayor_es_mejor"),
    ("redundancy_pct", "Redundancia (%)", "menor_es_mejor"),
    ("assertion_count", "Cantidad de Aserciones", "neutral"),
    ("test_count", "Cantidad de Tests", "neutral"),
]

# ── Colores ──────────────────────────────────────────────────────────────────
C_JUAMPI = "1F4E79"
C_MUTATION = "833C00"
C_HEADER = "404040"
C_ALT = "F2F2F2"
C_YES = "C6EFCE"
C_NO = "FFCCCC"
C_SIG = "FFF2CC"
C_TITLE = "D6E4F0"


def _border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def _hfill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _style_header(cell, color=C_HEADER):
    cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=9)
    cell.fill = _hfill(color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border()


def _style_cell(cell, fill=None, bold=False, align="center"):
    cell.font = Font(name="Arial", size=9, bold=bold)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = _border()
    if fill:
        cell.fill = _hfill(fill)


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _num(v):
    """Convierte a float si es posible, si no devuelve None (celda vacía o no numérica)."""
    if v is None:
        return None
    v = str(v).strip()
    if v == "":
        return None
    try:
        return float(v)
    except ValueError:
        return None


# ── Carga y deduplicación de metrics.csv ────────────────────────────────────
def load_metrics(path: Path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    # Deduplicar por (bug_id, class_under_test, algorithm, seed, java_version):
    # si hay corridas repetidas (--force), preferir la última fila cuyo
    # resultado no sea RUNNER ERROR; si todas son error, quedarse con la primera.
    dedup = {}
    for row in rows:
        key = (
            row["bug_id"].strip(),
            row["class_under_test"].strip(),
            row["algorithm"].strip(),
            row["seed"].strip(),
            row["java_version"].strip(),
        )
        if key not in dedup:
            dedup[key] = row
        else:
            prev = dedup[key]
            prev_is_error = prev.get("result", "").strip() == "RUNNER ERROR"
            curr_is_error = row.get("result", "").strip() == "RUNNER ERROR"
            if prev_is_error and not curr_is_error:
                dedup[key] = row
            elif not prev_is_error and not curr_is_error:
                dedup[key] = row  # se queda con la más reciente (última en el CSV)
    return list(dedup.values())


# ── Emparejamiento JUAMPI vs MUTATION ───────────────────────────────────────
def build_pairs(rows):
    """
    Devuelve dict: (bug_id, class_under_test, seed, java_version) ->
        {"JUAMPI": row, "MUTATION": row}
    Solo incluye combinaciones con AMBOS algoritmos presentes.
    """
    by_key = defaultdict(dict)
    for row in rows:
        key = (
            row["bug_id"].strip(),
            row["class_under_test"].strip(),
            row["seed"].strip(),
            row["java_version"].strip(),
        )
        algo = row["algorithm"].strip().upper()
        if algo in ("JUAMPI", "MUTATION"):
            by_key[key][algo] = row

    return {k: v for k, v in by_key.items() if "JUAMPI" in v and "MUTATION" in v}


# ── Vargha-Delaney A12 ───────────────────────────────────────────────────────
def vargha_delaney_a12(x: list, y: list) -> float:
    """
    A12 = P(valor aleatorio de x > valor aleatorio de y) + 0.5 * P(empate)
    A12 > 0.5  -> x tiende a ser mayor que y
    A12 < 0.5  -> x tiende a ser menor que y
    A12 = 0.5  -> no hay diferencia sistemática
    """
    m, n = len(x), len(y)
    if m == 0 or n == 0:
        return float("nan")
    wins = ties = 0
    for xi in x:
        for yi in y:
            if xi > yi:
                wins += 1
            elif xi == yi:
                ties += 1
    return (wins + 0.5 * ties) / (m * n)


def a12_magnitude(a12: float) -> str:
    if a12 != a12:  # NaN
        return "N/A"
    m = abs(a12 - 0.5)
    if m < 0.06:
        return "Insignificante"
    elif m < 0.14:
        return "Pequeño"
    elif m < 0.21:
        return "Mediano"
    return "Grande"


def a12_direction(a12: float, metric_type: str) -> str:
    if a12 != a12:
        return "N/A"
    if abs(a12 - 0.5) < 0.06:
        return "Sin diferencia"
    favors_juampi = a12 > 0.5
    if metric_type == "menor_es_mejor":
        favors_juampi = not favors_juampi
    return "Favorece JUAMPI" if favors_juampi else "Favorece MUTATION"


# ── Hoja: Vista Global ───────────────────────────────────────────────────────
def write_vista_global(wb, rows, pairs):
    ws = wb.create_sheet("Vista Global")
    ws.cell(row=1, column=1, value="VISTA GLOBAL — JUAMPI vs MUTATION").font = Font(bold=True, size=13)

    juampi_rows = [r for r in rows if r["algorithm"].strip().upper() == "JUAMPI"]
    mutation_rows = [r for r in rows if r["algorithm"].strip().upper() == "MUTATION"]

    def avg_metric(rowset, field):
        vals = [_num(r.get(field)) for r in rowset]
        vals = [v for v in vals if v is not None]
        return round(mean(vals), 2) if vals else None

    def detect_rate(rowset):
        vals = [1 if r.get("detected", "0").strip() == "1" else 0 for r in rowset]
        return round(sum(vals) / len(vals) * 100, 1) if vals else 0.0

    bugs_juampi = {(r["bug_id"], r["class_under_test"]) for r in juampi_rows}
    bugs_mutation = {(r["bug_id"], r["class_under_test"]) for r in mutation_rows}

    r = 3
    headers = ["Métrica", "JUAMPI", "MUTATION", "Nota"]
    for c, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=r, column=c, value=h))
    r += 1

    filas = [
        ("Total de corridas (bug x seed x java)", len(juampi_rows), len(mutation_rows), ""),
        ("Bugs/clases distintas evaluadas", len(bugs_juampi), len(bugs_mutation), ""),
        ("Pares comparables (mismo bug+seed+java)", len(pairs), len(pairs), "Usados en Estadísticos"),
        ("Tasa de detección de bugs (%)", detect_rate(juampi_rows), detect_rate(mutation_rows), "% BUG-REVEALING"),
    ]
    for m_key, m_label, _ in METRICAS:
        filas.append((f"Promedio {m_label}", avg_metric(juampi_rows, m_key), avg_metric(mutation_rows, m_key), ""))

    for row_data in filas:
        for c, v in enumerate(row_data, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            _style_cell(cell, fill=C_ALT if r % 2 == 0 else None, bold=(c == 1), align="left" if c in (1, 4) else "center")
        r += 1

    _autosize(ws, [42, 16, 16, 32])
    ws.freeze_panes = "A4"


# ── Hoja: Resumen (una fila por bug, promedio de seeds) ─────────────────────
def write_resumen(wb, pairs):
    ws = wb.create_sheet("Resumen")

    headers = ["bug_id", "class_under_test", "n_pares"]
    for m_key, m_label, _ in METRICAS:
        headers += [f"{m_label} (JUAMPI)", f"{m_label} (MUTATION)", f"{m_label} (Δ)"]
    headers += ["Detección JUAMPI (%)", "Detección MUTATION (%)"]

    for c, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=1, column=c, value=h))

    # Agrupar pares por bug_id + class_under_test
    by_bug = defaultdict(list)
    for (bug_id, class_under_test, seed, java_version), algos in pairs.items():
        by_bug[(bug_id, class_under_test)].append(algos)

    r = 2
    for (bug_id, class_under_test), pair_list in sorted(by_bug.items()):
        row_vals = [bug_id, class_under_test, len(pair_list)]
        for m_key, m_label, _ in METRICAS:
            j_vals = [_num(p["JUAMPI"].get(m_key)) for p in pair_list]
            m_vals = [_num(p["MUTATION"].get(m_key)) for p in pair_list]
            j_vals = [v for v in j_vals if v is not None]
            m_vals = [v for v in m_vals if v is not None]
            j_avg = round(mean(j_vals), 2) if j_vals else None
            m_avg = round(mean(m_vals), 2) if m_vals else None
            delta = round(j_avg - m_avg, 2) if (j_avg is not None and m_avg is not None) else None
            row_vals += [j_avg, m_avg, delta]

        j_det = [1 if p["JUAMPI"].get("detected", "0").strip() == "1" else 0 for p in pair_list]
        m_det = [1 if p["MUTATION"].get("detected", "0").strip() == "1" else 0 for p in pair_list]
        row_vals += [
            round(sum(j_det) / len(j_det) * 100, 1) if j_det else 0.0,
            round(sum(m_det) / len(m_det) * 100, 1) if m_det else 0.0,
        ]

        fill = C_ALT if r % 2 == 0 else None
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            _style_cell(cell, fill=fill, align="left" if c in (1, 2) else "center")
        r += 1

    widths = [42, 45, 10] + [16, 16, 12] * len(METRICAS) + [16, 16]
    _autosize(ws, widths)
    ws.freeze_panes = "A2"


# ── Hoja: Detección de Bugs ───────────────────────────────────────────────────
def write_deteccion(wb, pairs):
    ws = wb.create_sheet("Detección")

    headers = [
        "bug_id", "class_under_test", "seed", "java_version",
        "JUAMPI detecta?", "MUTATION detecta?", "Acuerdo",
    ]
    for c, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=1, column=c, value=h))

    r = 2
    for (bug_id, class_under_test, seed, java_version), algos in sorted(pairs.items()):
        j_det = algos["JUAMPI"].get("detected", "0").strip() == "1"
        m_det = algos["MUTATION"].get("detected", "0").strip() == "1"
        if j_det and m_det:
            acuerdo = "Ambos detectan"
        elif not j_det and not m_det:
            acuerdo = "Ninguno detecta"
        elif j_det:
            acuerdo = "Solo JUAMPI"
        else:
            acuerdo = "Solo MUTATION"

        row_vals = [
            bug_id, class_under_test, seed, java_version,
            "SI" if j_det else "NO", "SI" if m_det else "NO", acuerdo,
        ]
        fill = C_ALT if r % 2 == 0 else None
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 5:
                _style_cell(cell, fill=C_YES if v == "SI" else C_NO)
            elif c == 6:
                _style_cell(cell, fill=C_YES if v == "SI" else C_NO)
            else:
                _style_cell(cell, fill=fill, align="left" if c in (1, 2) else "center")
        r += 1

    _autosize(ws, [42, 45, 16, 12, 15, 17, 18])
    ws.freeze_panes = "A2"


# ── Hoja: Estadísticos ────────────────────────────────────────────────────────
def write_estadisticos(wb, pairs):
    ws = wb.create_sheet("Estadísticos")
    ws.cell(row=1, column=1, value="ANÁLISIS ESTADÍSTICO — Wilcoxon signed-rank + Vargha-Delaney A12").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="H0: no hay diferencia entre JUAMPI y MUTATION. Se rechaza si p < 0.05.").font = Font(italic=True, size=9)

    headers = [
        "Métrica", "N pares", "Media JUAMPI", "Media MUTATION",
        "Estadístico W", "p-valor", "¿Significativo? (p<0.05)",
        "A12", "Magnitud (A12)", "Dirección",
    ]
    r = 4
    for c, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=r, column=c, value=h))
    r += 1

    for m_key, m_label, m_type in METRICAS:
        j_vals, m_vals = [], []
        for algos in pairs.values():
            jv = _num(algos["JUAMPI"].get(m_key))
            mv = _num(algos["MUTATION"].get(m_key))
            if jv is not None and mv is not None:
                j_vals.append(jv)
                m_vals.append(mv)

        n = len(j_vals)
        if n < 2:
            row_vals = [m_label, n, "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "N/A", "Datos insuficientes"]
        else:
            j_mean = round(mean(j_vals), 2)
            m_mean = round(mean(m_vals), 2)

            diffs = [j - m for j, m in zip(j_vals, m_vals)]
            all_zero = all(d == 0 for d in diffs)

            if wilcoxon is None:
                stat_val, p_val = "N/A", "N/A"
                significativo = "scipy no disponible"
            elif all_zero:
                stat_val, p_val = 0.0, 1.0
                significativo = "NO"
            else:
                try:
                    stat_val, p_val = wilcoxon(j_vals, m_vals)
                    stat_val = round(float(stat_val), 3)
                    p_val = round(float(p_val), 4)
                    significativo = "SI" if p_val < 0.05 else "NO"
                except ValueError:
                    # ocurre si todas las diferencias son cero tras el ajuste interno
                    stat_val, p_val, significativo = 0.0, 1.0, "NO"

            a12 = vargha_delaney_a12(j_vals, m_vals)
            magnitud = a12_magnitude(a12)
            direccion = a12_direction(a12, m_type)

            row_vals = [
                m_label, n, j_mean, m_mean, stat_val, p_val, significativo,
                round(a12, 3) if a12 == a12 else "N/A", magnitud, direccion,
            ]

        fill = C_ALT if r % 2 == 0 else None
        for c, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 7 and v == "SI":
                _style_cell(cell, fill=C_SIG, bold=True)
            elif c == 1:
                _style_cell(cell, fill=fill, bold=True, align="left")
            else:
                _style_cell(cell, fill=fill)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Guía de magnitud A12 (Vargha & Delaney, 2000):").font = Font(italic=True, size=8)
    r += 1
    ws.cell(row=r, column=1, value="|A12 - 0.5| < 0.06: Insignificante | 0.06-0.14: Pequeño | 0.14-0.21: Mediano | ≥0.21: Grande").font = Font(italic=True, size=8)

    _autosize(ws, [30, 10, 14, 16, 14, 12, 20, 10, 16, 20])
    ws.freeze_panes = "A5"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--metrics", default=str(DEFAULT_METRICS))
    parser.add_argument("--out", default=str(DEFAULT_OUT))
    args = parser.parse_args()

    metrics_path = Path(args.metrics)
    out_path = Path(args.out)

    if not metrics_path.is_file():
        print(f"ERROR: no se encontró {metrics_path}")
        return

    if wilcoxon is None:
        print("[AVISO] scipy no está instalado — los p-valores saldrán como 'N/A'.")
        print("        Instalar con: pip install scipy --break-system-packages")

    print(f"Leyendo {metrics_path}...")
    rows = load_metrics(metrics_path)
    print(f"  {len(rows)} filas (tras deduplicar corridas repetidas)")

    pairs = build_pairs(rows)
    print(f"  {len(pairs)} pares JUAMPI-MUTATION comparables (mismo bug+seed+java_version)")

    wb = Workbook()
    wb.remove(wb.active)

    write_vista_global(wb, rows, pairs)
    write_resumen(wb, pairs)
    write_deteccion(wb, pairs)
    write_estadisticos(wb, pairs)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"\n  Guardado: {out_path}")
    print("\nHojas generadas:")
    print("  - Vista Global  : totales y promedios generales")
    print("  - Resumen       : comparación por bug (promedio de seeds)")
    print("  - Detección     : detalle BUG-REVEALING por corrida")
    print("  - Estadísticos  : Wilcoxon signed-rank + Vargha-Delaney A12")


if __name__ == "__main__":
    main()