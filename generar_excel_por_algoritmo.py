#!/usr/bin/env python3
"""
Paso 1: Lee un results.csv (salida de run_screened_bugs.py), parsea los logs
reales de EvoSuite y los .java de tests generados, y produce:

  - metrics.csv   (append mode, para combinar JUAMPI + MUTATION en un solo archivo)
  - resultados_<ALGORITMO>.xlsx  (una hoja por bug con el detalle de cada seed)

A diferencia de la versión con mocks, este script parsea directamente:
  1. logs/evosuite_fixed.log       -> coverage, mutation score, algoritmo, test count
  2. fixed/evosuite-tests/.../<Clase>_ESTest.java -> aserciones reales, redundancia real

Detección de algoritmo:
  Si el log contiene una línea "strategy: JUAMPI" -> JUAMPI
  Si no aparece esa línea -> MUTATION (el algoritmo nativo no la imprime)
  --algoritmo fuerza la etiqueta si preferís no confiar en el log.

Definición de redundancia (según operacionalización de variables de la tesis):
  Dos aserciones son redundantes si coinciden en:
    (tipo de aserción, método invocado, valor esperado)
  sin importar cómo se construyó el objeto ni en qué test aparecen.
  % Redundancia = aserciones redundantes / total aserciones x 100
  ("redundantes" = copias adicionales más allá de la primera aparición de cada clave única)

Uso:
    python3 generar_excel_por_algoritmo.py \
        --results ~/Tesis/gitbug-batch/results.csv \
        --out-dir ~/Tesis/gitbug-batch/ \
        --algoritmo MUTATION

    python3 generar_excel_por_algoritmo.py \
        --results ~/Tesis/gitbug-batch2/results.csv \
        --out-dir ~/Tesis/gitbug-batch2/ \
        --algoritmo JUAMPI
"""

import argparse
import csv
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Regex para parsear evosuite_fixed.log ──────────────────────────────────
STRATEGY_RE = re.compile(r'^strategy:\s*(\w+)', re.MULTILINE)
LINE_COV_RE = re.compile(r"Coverage of criterion LINE:\s*(\d+)%")
BRANCH_COV_RE = re.compile(r"Coverage of criterion BRANCH:\s*(\d+)%")
MUTSCORE_RE = re.compile(r"Resulting test suite's mutation score:\s*(\d+)%")
TESTCOUNT_RE = re.compile(r"Generated (\d+) tests? with total length")

# ── Regex para parsear aserciones en los .java generados ───────────────────
ASSERT_RE = re.compile(
    r'\b(assertEquals|assertTrue|assertFalse|assertNull|assertNotNull|'
    r'assertSame|assertNotSame|assertArrayEquals)\s*\((.*?)\);',
    re.DOTALL,
)
# variable receptora tipo "map0.isEmpty()" -> se quita "map0." para no contar
# la construcción del objeto como parte de la identidad de la aserción
VAR_PREFIX_RE = re.compile(r'^[a-zA-Z_][a-zA-Z0-9_]*?(\d+)\.')
TEST_METHOD_SPLIT_RE = re.compile(r'@Test\([^)]*\)\s*\n\s*public void (\w+)\(\)')


def parse_log(log_text: str, forced_algorithm: str = "") -> dict:
    """Extrae algoritmo, cobertura, mutation score y cantidad de tests del log."""
    m = {}

    if forced_algorithm:
        m["algorithm"] = forced_algorithm
    else:
        algo_match = STRATEGY_RE.search(log_text)
        m["algorithm"] = algo_match.group(1) if algo_match else "MUTATION"

    def grab(pattern):
        mm = pattern.search(log_text)
        return mm.group(1) if mm else ""

    m["line_coverage"] = grab(LINE_COV_RE)
    m["branch_coverage"] = grab(BRANCH_COV_RE)
    m["mutation_score_final"] = grab(MUTSCORE_RE)
    m["test_count"] = grab(TESTCOUNT_RE)
    return m


def split_top_level_commas(s: str):
    """Separa argumentos por comas de nivel superior, respetando paréntesis y strings."""
    parts = []
    depth = 0
    current = ""
    in_string = False
    escape = False
    for ch in s:
        if escape:
            current += ch
            escape = False
            continue
        if ch == "\\":
            current += ch
            escape = True
            continue
        if ch == '"':
            in_string = not in_string
            current += ch
            continue
        if not in_string:
            if ch in "([":
                depth += 1
            elif ch in ")]":
                depth -= 1
            elif ch == "," and depth == 0:
                parts.append(current.strip())
                current = ""
                continue
        current += ch
    if current.strip():
        parts.append(current.strip())
    return parts


def normalize_expr(expr: str) -> str:
    """
    Normaliza una expresión Java para comparar aserciones ignorando el nombre
    de la variable receptora (que depende de cómo EvoSuite construyó el objeto).
    Ej: "map0.isEmpty()" y "map1.isEmpty()" -> "isEmpty()"
    """
    expr = expr.strip()
    expr = re.sub(r"^\([^()]*\)\s*", "", expr)  # quita casts: (Property<?>) x0...
    expr = VAR_PREFIX_RE.sub("", expr)
    return expr


def parse_test_file(java_text: str):
    """
    Devuelve lista de (nombre_test, clave_asercion) donde
    clave_asercion = (tipo_assert, expresion_normalizada, valor_esperado)
    """
    results = []
    blocks = TEST_METHOD_SPLIT_RE.split(java_text)
    # blocks[0] = preámbulo; luego alterna (nombre_test, cuerpo)
    for i in range(1, len(blocks), 2):
        name = blocks[i]
        body = blocks[i + 1]
        for match in ASSERT_RE.finditer(body):
            assert_type = match.group(1)
            args = split_top_level_commas(match.group(2))

            if assert_type == "assertEquals" and len(args) >= 2:
                expected = args[0].strip()
                actual = normalize_expr(args[-1])
                key = (assert_type, actual, expected)
            elif assert_type in ("assertTrue", "assertFalse") and len(args) >= 1:
                actual = normalize_expr(args[0])
                expected = "true" if assert_type == "assertTrue" else "false"
                key = (assert_type, actual, expected)
            elif assert_type in ("assertNull", "assertNotNull") and len(args) >= 1:
                actual = normalize_expr(args[0])
                key = (assert_type, actual, "")
            else:
                actual = normalize_expr(args[-1]) if args else ""
                expected = args[0] if len(args) > 1 else ""
                key = (assert_type, actual, expected)

            results.append((name, key))
    return results


def find_test_file(work_root: Path, class_under_test: str) -> Path | None:
    """
    Busca el .java generado por EvoSuite para una clase dada, dentro de
    <work_root>/fixed/evosuite-tests/<paquete>/<Clase>_ESTest.java
    """
    class_name = class_under_test.rsplit(".", 1)[-1]
    pkg_path = class_under_test.rsplit(".", 1)[0].replace(".", "/") if "." in class_under_test else ""

    candidate = work_root / "fixed" / "evosuite-tests" / pkg_path / f"{class_name}_ESTest.java"
    if candidate.is_file():
        return candidate

    # Fallback: buscar en cualquier lado bajo evosuite-tests por si el paquete no calza
    tests_root = work_root / "fixed" / "evosuite-tests"
    if tests_root.is_dir():
        matches = list(tests_root.rglob(f"{class_name}_ESTest.java"))
        if matches:
            return matches[0]

    return None


def compute_assertion_metrics(java_path: Path):
    """Devuelve (total_aserciones, aserciones_redundantes, %redundancia)."""
    java_text = java_path.read_text(encoding="utf-8", errors="replace")
    assertions = parse_test_file(java_text)
    total = len(assertions)
    if total == 0:
        return 0, 0, 0.0

    keys = [k for _, k in assertions]
    counts = Counter(keys)
    redundant = sum(c - 1 for c in counts.values() if c > 1)
    pct = round(redundant / total * 100, 1)
    return total, redundant, pct


# ── Estilos Excel ────────────────────────────────────────────────────────────
C_HEADER = "1F4E79"
C_ALT = "F2F2F2"
C_YES = "C6EFCE"
C_NO = "FFCCCC"
C_WARN = "FFF2CC"


def _border():
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def _style_header(cell):
    cell.font = Font(bold=True, color="FFFFFF", name="Arial", size=9)
    cell.fill = PatternFill("solid", start_color=C_HEADER, end_color=C_HEADER)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border()


def _style_cell(cell, fill=None, bold=False):
    cell.font = Font(name="Arial", size=9, bold=bold)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _border()
    if fill:
        cell.fill = PatternFill("solid", start_color=fill, end_color=fill)


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_bug_sheet(wb: Workbook, bug_id: str, rows: list):
    # Excel no permite ciertos caracteres ni más de 31 chars en nombres de hoja
    safe_name = re.sub(r'[\\/*?:\[\]]', "_", bug_id)[:31]
    ws = wb.create_sheet(title=safe_name)

    headers = [
        "class_under_test", "seed", "java_version", "result", "detected",
        "line_cov(%)", "branch_cov(%)", "mutation_score(%)",
        "test_count", "assertion_count", "redundant_assertions", "redundancy(%)",
    ]
    for col, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=1, column=col, value=h))

    for r, row in enumerate(rows, start=2):
        values = [
            row["class_under_test"], row["seed"], row["java_version"], row["result"],
            "SI" if row["detected"] == "1" else "NO",
            row["line_coverage"], row["branch_coverage"], row["mutation_score_final"],
            row["test_count"], row["assertion_count"], row["redundant_assertion_count"],
            row["redundancy_pct"],
        ]
        fill = C_ALT if r % 2 == 0 else None
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 5:  # columna detected
                _style_cell(cell, fill=C_YES if v == "SI" else C_NO)
            else:
                _style_cell(cell, fill=fill)

    _autosize(ws, [42, 16, 12, 20, 10, 12, 14, 16, 12, 15, 18, 13])
    ws.freeze_panes = "A2"


def main():
    parser = argparse.ArgumentParser(description="Genera metrics.csv y Excel a partir de resultados reales de EvoSuite.")
    parser.add_argument("--results", required=True, help="Ruta al results.csv generado por run_screened_bugs.py")
    parser.add_argument("--out-dir", required=True, help="Carpeta donde guardar metrics.csv y el Excel")
    parser.add_argument("--algoritmo", default="", choices=["", "JUAMPI", "MUTATION"],
                         help="Fuerza la etiqueta de algoritmo en vez de detectarla del log")
    args = parser.parse_args()

    results_path = Path(args.results)
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    if not results_path.is_file():
        print(f"ERROR: no se encontró {results_path}")
        return

    with open(results_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows_in = [row for row in reader]

    print(f"Leyendo {results_path} ({len(rows_in)} filas)...")

    all_rows = []
    bugs_by_id = defaultdict(list)
    missing_log = 0
    missing_test_file = 0

    for row in rows_in:
        bug_id = row["bug_id"].strip()
        class_under_test = row["class_under_test"].strip()
        seed = row["seed"].strip()
        java_version = row["java_version"].strip()
        result = row["result"].strip()
        work_root = Path(row["work_root"].strip())
        logs_dir = Path(row["logs_dir"].strip())

        metrics = {
            "algorithm": args.algoritmo or "MUTATION",
            "line_coverage": "", "branch_coverage": "",
            "mutation_score_final": "", "test_count": "",
        }
        assertion_count = ""
        redundant_count = ""
        redundancy_pct = ""

        log_path = logs_dir / "evosuite_fixed.log"
        if log_path.is_file():
            log_text = log_path.read_text(encoding="utf-8", errors="replace")
            metrics = parse_log(log_text, forced_algorithm=args.algoritmo)
        else:
            missing_log += 1

        test_file = find_test_file(work_root, class_under_test)
        if test_file is not None:
            total, redundant, pct = compute_assertion_metrics(test_file)
            assertion_count, redundant_count, redundancy_pct = total, redundant, pct
        else:
            missing_test_file += 1

        out_row = {
            "bug_id": bug_id,
            "class_under_test": class_under_test,
            "algorithm": metrics["algorithm"],
            "seed": seed,
            "java_version": java_version,
            "result": result,
            "detected": "1" if result == "BUG-REVEALING" else "0",
            "line_coverage": metrics["line_coverage"],
            "branch_coverage": metrics["branch_coverage"],
            "mutation_score_final": metrics["mutation_score_final"],
            "test_count": metrics["test_count"],
            "assertion_count": assertion_count,
            "redundant_assertion_count": redundant_count,
            "redundancy_pct": redundancy_pct,
        }
        all_rows.append(out_row)
        bugs_by_id[bug_id].append(out_row)

    if missing_log:
        print(f"  [AVISO] {missing_log} filas sin evosuite_fixed.log (corridas fallidas u otras rutas)")
    if missing_test_file:
        print(f"  [AVISO] {missing_test_file} filas sin archivo de test _ESTest.java encontrado")

    # ── Guardar metrics.csv (append) ────────────────────────────────────────
    fieldnames = [
        "bug_id", "class_under_test", "algorithm", "seed", "java_version",
        "result", "detected", "line_coverage", "branch_coverage",
        "mutation_score_final", "test_count", "assertion_count",
        "redundant_assertion_count", "redundancy_pct",
    ]
    metrics_csv = out_dir / "metrics.csv"
    file_exists = metrics_csv.is_file()
    with open(metrics_csv, "a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        if not file_exists:
            w.writeheader()
        w.writerows(all_rows)
    modo = "Agregado a" if file_exists else "Guardado"
    print(f"  {modo}: {metrics_csv}  ({len(all_rows)} filas nuevas)")

    # ── Generar Excel por algoritmo, una hoja por bug ───────────────────────
    algo_label = args.algoritmo or (all_rows[0]["algorithm"] if all_rows else "RESULTADOS")
    wb = Workbook()
    wb.remove(wb.active)
    for bug_id, rows in sorted(bugs_by_id.items()):
        write_bug_sheet(wb, bug_id, rows)

    out_xlsx = out_dir / f"resultados_{algo_label}.xlsx"
    wb.save(out_xlsx)
    print(f"  Guardado: {out_xlsx}  ({len(bugs_by_id)} hojas, 1 por bug)")


if __name__ == "__main__":
    main()